"""Write the priced bill: the owner's own Excel with the rate column filled, plus an "Agent report" sheet."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .pipeline import PRICED, CHECK, ASK

FILL = {PRICED: None, CHECK: PatternFill("solid", fgColor="FFF2CC"), ASK: PatternFill("solid", fgColor="F8CBAD")}
LABEL = {PRICED: "Priced / تم التسعير", CHECK: "Please check / راجع", ASK: "Needs you / يحتاج ردك"}
HEAD = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(vertical="top", wrap_text=True)
RTL = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)


def write_priced_bill(out, path_out, price_list_name="", lang="en"):
    bill, results, summ = out["bill"], out["results"], out["summary"]
    wb = openpyxl.load_workbook(bill["path"])
    ws = wb[bill["sheet"]]
    rate_col, total_col, qty_col = bill["cols"].get("rate"), bill["cols"].get("total"), bill["cols"]["qty"]
    for r in results:
        if rate_col:
            c = ws.cell(r["row"], rate_col)
            c.value = r["unit_price"] if r["status"] in (PRICED, CHECK) else None
            c.number_format = "#,##0.00"
            if FILL[r["status"]]:
                c.fill = FILL[r["status"]]
        if total_col and r["unit_price"] is not None:
            t = ws.cell(r["row"], total_col)
            if not (isinstance(t.value, str) and t.value.startswith("=")):
                t.value = "=%s%d*%s%d" % (_col(qty_col), r["row"], _col(rate_col), r["row"])

    rep = wb.create_sheet("Agent report")
    rows = [
        ("Bill", bill["path"].replace("\\", "/").split("/")[-1]),
        ("Priced on", datetime.date.today().isoformat()),
        ("Price list", price_list_name),
        ("Lines", summ["lines"]),
        ("Priced", summ["priced"]),
        ("Please check (priced, confirm)", summ["check"]),
        ("Needs you (no price yet)", summ["ask"]),
        ("Margin added", "%.0f%%" % (summ["margin"] * 100)),
        ("VAT 14% added", "yes" if summ["vat"] else "no"),
        ("AI calls / tokens", "%d / %d %s" % (summ["ai_calls"], summ["ai_tokens"], summ["ai_model"])),
        ("Total of priced lines (EGP)", summ["total_priced"]),
    ]
    for i, (k, v) in enumerate(rows, 1):
        rep.cell(i, 2, k).font = Font(bold=True)
        rep.cell(i, 3, v).alignment = Alignment(horizontal="left")
    rep.cell(len(rows), 3).number_format = "#,##0"
    hdr = len(rows) + 2
    heads = ["Line", "Section", "Description", "Unit", "Quantity", "Status", "Price-list code", "Price-list item",
             "Unit price (EGP)", "Total (EGP)", "Source", "Notes / question"]
    widths = [7, 22, 60, 7, 10, 20, 11, 40, 14, 15, 40, 60]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = rep.cell(hdr, i, h)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        rep.column_dimensions[_col(i)].width = w
    for j, r in enumerate(results, hdr + 1):
        vals = [r["key"], r["section"], r["desc"][:300], r["unit_raw"], r["qty"], LABEL[r["status"]], r["code"],
                r["item"], r["unit_price"], r["total"], r["source"], " ".join(n[lang] for n in r["notes"])]
        for i, v in enumerate(vals, 1):
            c = rep.cell(j, i, v)
            c.border = BOX
            c.alignment = RTL if i in (2, 3, 4, 8) else WRAP
            if i in (5, 9, 10):
                c.number_format = "#,##0.00" if i != 5 else "#,##0.##"
        if FILL[r["status"]]:
            rep.cell(j, 6).fill = FILL[r["status"]]
    rep.freeze_panes = rep.cell(hdr + 1, 4)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path_out)
    return path_out


def _col(i):
    s = ""
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s
