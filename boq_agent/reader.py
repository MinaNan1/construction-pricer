"""Read a bill of quantities (مقايسة) from Excel: find the header row, the columns, section titles and item lines."""
import openpyxl
from .text import normalize
from .units import parse_unit

# Header words per column role (compared after normalize()).
_ROLES = {
    "no": ["م", "رقم", "رقم البند", "مسلسل", "بند", "item", "no", "item no", "s no", "ser"],
    "desc": ["بيان الاعمال", "بيان", "وصف البند", "الوصف", "اسم البند", "بند الاعمال", "البند", "description", "item description", "works"],
    "unit": ["الوحده", "وحده", "unit"],
    "qty": ["الكميه", "كميه", "qty", "quantity"],
    "rate": ["الفئه", "فئه", "سعر الوحده", "السعر", "unit price", "rate", "price"],
    "total": ["الاجمالي", "اجمالي", "الاجمالى", "القيمه", "total", "amount"],
}
_TOTAL_WORDS = ("الاجمالي", "اجمالي", "المجموع", "total", "grand total", "subtotal")
_ROLES = {role: [normalize(w) for w in words] for role, words in _ROLES.items()}   # same cleanup as the cells
_TOTAL_WORDS = tuple(normalize(w) for w in _TOTAL_WORDS)


def _role(cell_text):
    n = normalize(cell_text)
    if not n:
        return None
    for role, words in _ROLES.items():          # exact words first ("م", "بند" = item number)
        if n in words:
            return role
    for role in ("desc", "rate", "total", "qty", "unit"):   # then "contains", e.g. "الفئة (جنيه)"
        if any(w in n for w in _ROLES[role] if len(w) > 2):
            return role
    return None


def _number(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def find_header(ws, max_rows=40):
    """Row number and {role: column} of the header row, or raises ValueError."""
    best = None
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cols = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                role = _role(v)
                if role and role not in cols:
                    cols[role] = c
        score = sum(k in cols for k in ("desc", "unit", "qty"))
        if score == 3:
            return r, cols
        if score >= 2 and (best is None or score > best[2]):
            best = (r, cols, score)
    if best and "desc" in best[1]:
        return best[0], best[1]
    raise ValueError("Could not find the header row (description / unit / quantity) in sheet '%s'" % ws.title)


def read_bill(path, sheet=None):
    """Returns dict(path, sheet, header_row, cols, lines). Each line: key, no, section, desc, unit_raw, unit, qty, row."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header_row, cols = find_header(ws)
    lines, section, sec_i = [], "", 0
    for r in range(header_row + 1, ws.max_row + 1):
        get = lambda role: ws.cell(r, cols[role]).value if role in cols else None
        desc = get("desc")
        desc = str(desc).strip() if desc is not None else ""
        qty = _number(get("qty"))
        unit_raw = get("unit")
        if not desc and qty is None:
            continue
        if normalize(desc) in _TOTAL_WORDS or any(normalize(desc).startswith(w) for w in _TOTAL_WORDS):
            if qty is None:
                continue
        if qty is None and not unit_raw:
            section, sec_i = desc, sec_i + 1          # a section title such as "ثانياً: أعمال الخرسانات"
            continue
        no = get("no")
        lines.append({
            "key": "%d.%s" % (sec_i, no if no is not None else len(lines) + 1),
            "no": no, "section": section, "desc": desc,
            "unit_raw": unit_raw, "unit": parse_unit(unit_raw), "qty": qty, "row": r,
        })
    return {"path": path, "sheet": ws.title, "header_row": header_row, "cols": cols, "lines": lines}
