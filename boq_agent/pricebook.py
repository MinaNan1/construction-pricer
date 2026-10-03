"""The price book: the approved price list workbook, recalculated in Python (pycel) so input changes
(e.g. a new steel price) update every item exactly the way Excel would."""
import datetime
import json
import os
import openpyxl
from pycel import ExcelCompiler
from .units import parse_unit
from .text import tokens

SHEET = "Price list"
COL = {"code": 1, "ar": 2, "en": 3, "unit": 4, "type": 5, "material": 6, "note": 16, "price": 15}


class PriceBook:
    def __init__(self, path, changes_path=None):
        """changes_path: JSON log of price changes made in the app; applied on top of the workbook,
        which itself is never modified."""
        self.path = path
        self.changes_path = changes_path
        self.xl = ExcelCompiler(filename=path)
        ws = openpyxl.load_workbook(path)[SHEET]
        self.items = []
        r = 2
        while ws.cell(r, COL["code"]).value:
            v = lambda k: ws.cell(r, COL[k]).value
            item = {
                "code": v("code"), "name_ar": v("ar"), "name_en": v("en"),
                "unit_label": v("unit"), "unit": parse_unit(v("unit")),
                "type": v("type"), "material": v("material"), "note": v("note") or "", "row": r,
            }
            item["tokens_ar"] = tokens(item["name_ar"])
            item["tokens_en"] = tokens(item["name_en"])
            self.items.append(item)
            r += 1
        self.by_code = {i["code"]: i for i in self.items}
        for c in self.changes():
            self.xl.set_value(c["address"], c["value"])

    def price(self, code):
        """Unit price of one item in EGP, recalculated from the current inputs."""
        return float(self.xl.evaluate("%s!O%d" % (SHEET, self.by_code[code]["row"])))

    def source(self, code):
        if code.startswith("CW"):
            return "Price list %s (Egyptian market cost model, July 2026)" % code
        return "Price list %s (Ministry of Housing bulletin Aug 2026 + rates)" % code

    def inputs(self):
        """Prices the owner may update: list of dicts (en, ar, address, value, unit)."""
        wb = openpyxl.load_workbook(self.path)
        out = []
        ws = wb["Inputs"]
        r = 4
        while ws.cell(r, 1).value:                                  # materials, EGP per ton
            out.append({"en": ws.cell(r, 1).value, "ar": ws.cell(r, 2).value, "address": "Inputs!C%d" % r,
                        "unit": "EGP/ton"})
            r += 1
        r += 2
        while ws.cell(r, 1).value:                                  # rates
            out.append({"en": ws.cell(r, 1).value, "ar": ws.cell(r, 2).value, "address": "Inputs!C%d" % r,
                        "unit": ws.cell(r, 5).value})
            r += 1
        ws = wb["Construction"]
        r = 4
        while ws.cell(r, 1).value and ws.cell(r, 5).value is not None:   # construction supplier prices
            out.append({"en": ws.cell(r, 2).value, "ar": ws.cell(r, 3).value, "address": "Construction!E%d" % r,
                        "unit": ws.cell(r, 4).value})
            r += 1
        for i in out:
            i["value"] = float(self.xl.evaluate(i["address"]))
        return out

    def changes(self):
        if self.changes_path and os.path.exists(self.changes_path):
            with open(self.changes_path, encoding="utf-8") as f:
                return json.load(f)
        return []

    def change_input(self, address, value, name=""):
        """Change an input and keep it in the change log so it survives a restart."""
        log = [c for c in self.changes() if c["address"] != address]
        log.append({"address": address, "value": value, "name": name, "date": datetime.date.today().isoformat()})
        if self.changes_path:
            with open(self.changes_path, "w", encoding="utf-8") as f:
                json.dump(log, f, ensure_ascii=False, indent=1)
        self.set_input(address, value)

    def set_input(self, address, value):
        """Change one input cell, e.g. set_input("Inputs!C4", 60000). Every price that uses it follows."""
        self.xl.set_value(address, value)
