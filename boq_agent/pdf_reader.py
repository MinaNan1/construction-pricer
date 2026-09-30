"""PDF bills: find the pages that hold the bill of quantities, have Gemini read only those pages into lines,
and write them to an Excel bill that the normal pipeline prices."""
import json
import re
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from .text import normalize

PDF_MODELS = ["gemini-3.1-flash-lite", "gemini-flash-lite-latest"]
MAX_PAGES = 15
_BOQ_WORDS = ("الكميه", "الوحده", "الفيه", "بيان الاعمال", "quantity", "unit", "rate")

SYSTEM = (
    "You read a bill of quantities (جدول الكميات / مقايسة) from tender pages. Return every priced line in order: "
    "no (the item number as printed), section (the section title it is under, e.g. \"أولاً: الأعمال المدنية\"), "
    "desc (the full item description, copied word for word; join text broken across lines; fix only letters that "
    "are clearly broken by the PDF), unit (as printed in the unit column), qty (a number). Skip headers, totals and "
    "blank price columns. Do not invent or translate anything. JSON only."
)
SCHEMA = {
    "type": "object",
    "properties": {"lines": {"type": "array", "items": {"type": "object", "properties": {
        "no": {"type": "string"}, "section": {"type": "string"}, "desc": {"type": "string"},
        "unit": {"type": "string"}, "qty": {"type": "number"}},
        "required": ["no", "section", "desc", "unit", "qty"]}}},
    "required": ["lines"],
}


def bill_pages(pdf_path):
    """Pages (0-based) that look like the bill table. Scanned PDFs without text: the first MAX_PAGES pages."""
    doc = pymupdf.open(pdf_path)
    hits, has_text = [], False
    for i, page in enumerate(doc):
        t = normalize(page.get_text())
        has_text = has_text or len(t) > 50
        if sum(w in t for w in _BOQ_WORDS) >= 2 and re.search(r"\d", t):
            hits.append(i)
    if not has_text:
        return list(range(min(len(doc), MAX_PAGES)))
    return hits[:MAX_PAGES]


def read_pdf_bill(pdf_path, resolver):
    """Returns (lines, pages) using Gemini on the bill pages only. Records calls/tokens on the resolver."""
    pages = bill_pages(pdf_path)
    if not pages or not resolver.gemini_key:
        return [], pages
    # pages go as pictures: PDF text layers often hold broken Arabic letters (swapped lam-alef), pictures don't
    src = pymupdf.open(pdf_path)
    images = [src[p].get_pixmap(dpi=130).tobytes("png") for p in pages]
    from google import genai
    from google.genai import types
    client = genai.Client(api_key=resolver.gemini_key)
    cfg = types.GenerateContentConfig(system_instruction=SYSTEM, temperature=0,
                                      response_mime_type="application/json", response_schema=SCHEMA)
    for model in PDF_MODELS:
        try:
            r = client.models.generate_content(
                model=model, config=cfg,
                contents=[types.Part.from_bytes(data=img, mime_type="image/png") for img in images]
                + ["Read the bill of quantities on these pages."])
            resolver.calls += 1
            u = r.usage_metadata
            resolver.tokens += (u.prompt_token_count or 0) + (u.candidates_token_count or 0)
            return json.loads(r.text)["lines"], pages
        except Exception as e:
            resolver.errors.append("pdf %s: %s" % (model, str(e)[:120]))
    return [], pages


def write_bill_xlsx(lines, path, title=""):
    """The standard bill layout the reader understands: م | بيان الأعمال | الوحدة | الكمية | الفئة | الإجمالي."""
    wb = Workbook()
    ws = wb.active
    ws.title = "المقايسة"
    ws.sheet_view.rightToLeft = True
    for col, w in zip("ABCDEF", [6, 70, 10, 11, 12, 15]):
        ws.column_dimensions[col].width = w
    ws["B1"] = title
    ws["B1"].font = Font(bold=True)
    for i, h in enumerate(["م", "بيان الأعمال", "الوحدة", "الكمية", "الفئة", "الإجمالي"], 1):
        ws.cell(3, i, h).font = Font(bold=True)
    r, section = 4, None
    for ln in lines:
        if ln.get("section") and ln["section"] != section:
            section = ln["section"]
            ws.cell(r, 2, section).font = Font(bold=True)
            r += 1
        ws.cell(r, 1, ln.get("no"))
        ws.cell(r, 2, ln.get("desc")).alignment = Alignment(wrap_text=True, horizontal="right", readingOrder=2)
        ws.cell(r, 3, ln.get("unit"))
        ws.cell(r, 4, ln.get("qty"))
        ws.cell(r, 6, '=IF(E{0}="","",D{0}*E{0})'.format(r))
        r += 1
    wb.save(path)
    return path
