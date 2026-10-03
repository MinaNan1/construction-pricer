"""Accuracy test: price the demo bills and compare with the answer keys.
Bill 1: correct item and price per line vs an engineer's priced bill. Bill 2: no confident wrong prices."""
import openpyxl
from boq_agent.pricebook import PriceBook
from boq_agent.pipeline import price_bill, PRICED, CHECK
from boq_agent.ai import Resolver

PRICE_LIST = "Price list - steel and metal works.xlsx"
BILL1 = "demo bills/Bill 1 - residential building.xlsx"
KEY1 = "demo bills/Bill 1 - answer key.xlsx"
BILL2 = "demo bills/Bill 2 - port fence and gates (SCZone tender 5 of 2020-2021).xlsx"
KEY2 = "demo bills/Bill 2 - expected behaviour.xlsx"


def bill1(pb):
    key = {}
    for row in openpyxl.load_workbook(KEY1).active.iter_rows(min_row=4, values_only=True):
        if isinstance(row[0], int):
            key[row[0]] = {"code": row[1], "qty": row[3], "price": row[4]}
    out = price_bill(BILL1, pb, resolver=Resolver())
    auto = [r for r in out["results"] if r["status"] in (PRICED, CHECK)]
    wrong = [r["key"] for r in auto if r["code"] != key[r["no"]]["code"]]
    top1 = sum(r["candidates"] and r["candidates"][0]["code"] == key[r["no"]]["code"] for r in out["results"])
    priced_err = [abs(r["unit_price"] - key[r["no"]]["price"]) / key[r["no"]]["price"] for r in auto]
    key_total = sum(k["qty"] * k["price"] for k in key.values())
    s = out["summary"]
    print("Bill 1 (%d lines): %d priced without asking (%d of them by the AI), %d of them wrong; "
          "right item ranked first by the code on %d/%d lines. AI: %d call(s), %d tokens."
          % (len(out["results"]), len(auto), s["by_ai"], len(wrong), top1, len(out["results"]), s["ai_calls"], s["ai_tokens"]))
    if priced_err:
        print("  Price error on priced lines: max %.2f%%" % (100 * max(priced_err)))
    print("  Priced total %s of %s EGP (%.0f%% of the bill value); the rest waits for AI/owner."
          % (format(out["summary"]["total_priced"], ",.0f"), format(key_total, ",.0f"),
             100 * out["summary"]["total_priced"] / key_total))
    return out


def bill2(pb):
    exp = {}
    for row in openpyxl.load_workbook(KEY2).active.iter_rows(min_row=4, values_only=True):
        if row[0] and row[0] != "Parsing":
            sec, no = row[0].split()
            exp[("1." if sec == "Civil" else "2.") + no] = (row[1], row[2])
    out = price_bill(BILL2, pb, resolver=Resolver())
    bad = [r["key"] for r in out["results"] if r["status"] == PRICED and exp[r["key"]][0].startswith("Ask")]
    wrong_code = [r["key"] for r in out["results"] if r["code"] and exp[r["key"]][1] and r["code"] != exp[r["key"]][1]]
    to_check = [r["key"] for r in out["results"] if r["status"] == CHECK and exp[r["key"]][0].startswith("Ask")]
    conflict = [r["key"] for r in out["results"] if any(n["kind"] == "conflict" for n in r["notes"])]
    s = out["summary"]
    print("Bill 2 (%d lines): %d priced, %d priced-but-check, %d asked. Wrong confident prices: %s. "
          "Wrong item picked: %s. Should-ask lines priced but flagged for checking: %s. Unit conflicts caught: %s. "
          "AI: %d call(s), %d tokens."
          % (len(out["results"]), s["priced"], s["check"], s["ask"], bad or "none", wrong_code or "none",
             to_check or "none", conflict or "none", s["ai_calls"], s["ai_tokens"]))
    return out


if __name__ == "__main__":
    pb = PriceBook(PRICE_LIST)
    bill1(pb)
    bill2(pb)
